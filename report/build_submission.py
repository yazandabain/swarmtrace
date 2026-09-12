"""Render the editable manuscript using the supplied Apart template's type and page style.

The DOCX is built from the original template package. PDF uses its embedded fonts,
Letter page geometry, margins, heading sizes, and bordered title/abstract arrangement.
No office suite is required. Both are rendered from the same Markdown blocks.
"""

import hashlib
import html
import json
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
TEMPLATE = ROOT / "Copy of Apart Research hackathon submission template.docx"
FONT = "Old Standard TT"


def blocks(text):
    lines = text.splitlines()
    i = 0
    result = []
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            result.append(("break", None))
        elif line.startswith("## "):
            result.append(("heading", line[3:]))
        elif line.startswith("### "):
            result.append(("subheading", line[4:]))
        elif line.startswith("!["):
            result.append(("image", re.search(r"\]\((.*?)\)", line).group(1)))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", v) for v in row):
                    rows.append(row)
                i += 1
            result.append(("table", rows))
            continue
        else:
            paragraph = [line]
            while i + 1 < len(lines) and lines[i + 1].strip():
                i += 1
                paragraph.append(lines[i].strip())
            value = " ".join(paragraph)
            result.append(
                ("caption" if re.match(r"(Figure|Table) \d+\.", value) else "body", value)
            )
        i += 1
    return result


def inline(text):
    text = html.escape(text)
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    # Plain source URLs remain visible and are also clickable.
    text = re.sub(r"(https://[^\s<]+)", r'<link href="\1" color="#21658c">\1</link>', text)
    return text


def set_text(paragraph, text, size=11, bold=False):
    for i, part in enumerate(re.split(r"(\*\*.*?\*\*)", text)):
        run = paragraph.add_run(part.strip("*").replace("`", ""))
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold or (i % 2 == 1)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(5.5)
    return paragraph


def docx_render(title, author, affiliation, abstract, content):
    doc = Document(TEMPLATE)
    original_header = deepcopy(doc.tables[0]._tbl)
    # Preserve the supplied section properties and package-level styles/fonts.
    for child in list(doc._element.body):
        if child.tag != qn("w:sectPr"):
            doc._element.body.remove(child)
    doc._element.body.insert(0, original_header)
    header = doc.tables[0]
    for row in header.rows:
        for cell in row.cells:
            for child in list(cell._tc):
                if child.tag != qn("w:tcPr"):
                    cell._tc.remove(child)
    p = header.cell(0, 0).add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_text(p, title, 20, bold=True)
    cell = header.cell(1, 0)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_text(p, author + ("\n" + affiliation if affiliation else ""), 11)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_text(p, "With Apart Research", 11)
    set_text(cell.add_paragraph(), "Abstract", 12, bold=True)
    set_text(cell.add_paragraph(), abstract)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    for kind, value in content:
        if kind == "break":
            doc.add_page_break()
        elif kind in ["heading", "subheading"]:
            p = doc.add_paragraph(style="Heading 2" if kind == "heading" else "Heading 3")
            set_text(p, value, 14 if kind == "heading" else 13, bold=True)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.keep_with_next = True
        elif kind == "image":
            doc.add_picture(str((REPORT / value).resolve()), width=Inches(6.5))
            p = doc.paragraphs[-1]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
        elif kind == "table":
            table = doc.add_table(rows=len(value), cols=len(value[0]))
            table.autofit = False
            for i, row in enumerate(value):
                for j, v in enumerate(row):
                    p = table.cell(i, j).paragraphs[0]
                    set_text(p, v, 9, bold=i == 0)
                    p.paragraph_format.space_after = Pt(3)
                trpr = table.rows[i]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
            border = OxmlElement("w:tblBorders")
            for edge in ["top", "bottom", "insideH"]:
                tag = OxmlElement("w:" + edge)
                tag.set(qn("w:val"), "single")
                tag.set(qn("w:sz"), "4")
                tag.set(qn("w:color"), "BBBBBB")
                border.append(tag)
            table._tbl.tblPr.append(border)
        else:
            set_text(doc.add_paragraph(), value, 9 if kind == "caption" else 11)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_text(footer, "", 8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    dest = REPORT / "swarmtrace_paper.docx"
    doc.save(dest)
    return dest


def pdf_render(title, author, affiliation, abstract, content):
    font_dir = Path("/tmp/swarmtrace-template-fonts")
    font_dir.mkdir(exist_ok=True)
    with ZipFile(TEMPLATE) as z:
        for style in ["regular", "bold", "italic"]:
            data = z.read(f"word/fonts/OldStandardTT-{style}.ttf")
            p = font_dir / f"{style}.ttf"
            p.write_bytes(data)
            pdfmetrics.registerFont(TTFont("Apart-" + style, str(p)))
    pdfmetrics.registerFontFamily(
        "Apart-regular",
        normal="Apart-regular",
        bold="Apart-bold",
        italic="Apart-italic",
        boldItalic="Apart-bold",
    )
    body = ParagraphStyle(
        "body",
        fontName="Apart-regular",
        fontSize=11,
        leading=12.65,
        spaceAfter=5.5,
        splitLongWords=True,
    )
    styles = {
        "body": body,
        "caption": ParagraphStyle("caption", parent=body, fontSize=9, leading=10.5, spaceAfter=7),
        "heading": ParagraphStyle(
            "heading",
            parent=body,
            fontName="Apart-bold",
            fontSize=14,
            leading=21,
            spaceBefore=10,
            spaceAfter=0,
            keepWithNext=True,
        ),
        "subheading": ParagraphStyle(
            "subheading",
            parent=body,
            fontName="Apart-bold",
            fontSize=13,
            leading=19.5,
            spaceBefore=10,
            spaceAfter=0,
            keepWithNext=True,
        ),
    }
    title_style = ParagraphStyle(
        "title",
        parent=body,
        fontName="Apart-bold",
        fontSize=20,
        leading=22,
        alignment=1,
        spaceAfter=8,
        spaceBefore=6,
    )
    center = ParagraphStyle("author", parent=body, alignment=1, spaceAfter=5)
    title_p = Paragraph(inline(title), title_style)
    header_body = [
        Paragraph(inline(author) + ("<br/>" + inline(affiliation) if affiliation else ""), center),
        Paragraph("With Apart Research", center),
        Paragraph("<b>Abstract</b>", body),
        Paragraph(inline(abstract), body),
    ]
    header = Table([[title_p], [header_body]], colWidths=[468])
    header.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 1, colors.black),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.black),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story = [header, Spacer(1, 5)]
    for kind, value in content:
        if kind == "break":
            story.append(PageBreak())
        elif kind == "image":
            path = (REPORT / value).resolve()
            image = Image(
                str(path),
                width=468,
                height=468 * Image(str(path)).imageHeight / Image(str(path)).imageWidth,
            )
            image.keepWithNext = True
            story.append(image)
        elif kind == "table":
            cellstyle = ParagraphStyle("cell", parent=body, fontSize=9, leading=10.5, spaceAfter=0)
            cells = [
                [
                    Paragraph(("<b>" + inline(v) + "</b>") if i == 0 else inline(v), cellstyle)
                    for v in row
                ]
                for i, row in enumerate(value)
            ]
            if value[0][0] == "Window":
                widths = [43, 75, 80, 93, 94, 83]
            elif len(value[0]) == 7:
                widths = [28, 106, 44, 65, 75, 75, 75]
            else:
                widths = [35, 74, 92, 92, 92, 83]
            table = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("LINEABOVE", (0, 0), (-1, 0), 0.7, colors.black),
                        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                        ("LINEBELOW", (0, -1), (-1, -1), 0.7, colors.black),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 3),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.extend([table, Spacer(1, 5)])
        else:
            story.append(Paragraph(inline(value), styles[kind]))

    def footer(canvas, document):
        canvas.setFont("Apart-regular", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawRightString(540, 40, str(document.page))

    dest = REPORT / "swarmtrace_paper.pdf"
    doc = SimpleDocTemplate(
        str(dest),
        pagesize=(612, 792),
        leftMargin=72,
        rightMargin=72,
        topMargin=72,
        bottomMargin=72,
        title=title,
        author=author,
        allowSplitting=1,
        invariant=1,
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return dest


def main():
    source = (REPORT / "manuscript.md").read_text()
    assert "\u2014" not in source, "No em dashes in the manuscript"
    for unfinished in [
        "[Affiliation to confirm]",
        "pending author",
        "Draft for author review",
        "The author must personally",
    ]:
        assert unfinished not in source, f"Resolve final-paper text: {unfinished}"
    title = source.splitlines()[0][2:]
    author = re.search(r"^Author: (.*)$", source, re.MULTILINE).group(1)
    affiliation_match = re.search(r"^Affiliation:[ \t]*(.*)$", source, re.MULTILINE)
    affiliation = affiliation_match.group(1) if affiliation_match else ""
    abstract = source.split("## Abstract\n\n")[1].split("\n\n## 1.")[0]
    assert 150 <= len(abstract.split()) <= 250
    content = blocks("## 1." + source.split("\n\n## 1.", 1)[1])
    docx_path = docx_render(title, author, affiliation, abstract, content)
    pdf_path = pdf_render(title, author, affiliation, abstract, content)
    pdf = pymupdf.open(pdf_path)
    review = []
    preview = REPORT / "preview"
    preview.mkdir(exist_ok=True)
    for old in preview.glob("page-*.png"):
        old.unlink()
    for i, page in enumerate(pdf):
        text = page.get_text()
        page.get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(preview / f"page-{i + 1}.png")
        review.append(
            {
                "page": i + 1,
                "words": len(text.split()),
                "first_text": text[:140],
                "last_text": text[-120:],
            }
        )
    (REPORT / "render_validation.json").write_text(
        json.dumps(
            {
                "abstract_words": len(abstract.split()),
                "total_pdf_pages": len(pdf),
                "main_pdf_pages": next(
                    i for i, page in enumerate(pdf) if "References" in page.get_text().splitlines()
                ),
                "pages": review,
                "template": TEMPLATE.name,
                "template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
                "manuscript_sha256": hashlib.sha256(source.encode()).hexdigest(),
                "pdf_sha256": hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
                "docx_sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest(),
                "font": "Embedded Old Standard TT from template",
                "page_geometry": "US Letter, one-inch margins",
                "rendering": "DOCX from native package; PDF from shared manuscript and template font/style",
            },
            indent=2,
        )
        + "\n"
    )
    print(json.dumps(review, indent=2))
    print(docx_path)
    print(pdf_path)


if __name__ == "__main__":
    main()
